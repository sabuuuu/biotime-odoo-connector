"""Generate an Excel preview of the hr.attendance records the sync would write to Odoo.

Reads ZKBioTime only (plus Odoo employees if reachable, read-only); nothing is written.
"""
import os
import sys
from collections import Counter
from datetime import datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

import zk_odoo_sync as zk

RAW_COLUMNS = [
    ("id", "ID BioTime"),
    ("emp_code", "Numéro d'employé"),
    ("first_name", "Prénom"),
    ("last_name", "Nom"),
    ("department", "Département"),
    ("punch_time", "Temps de pointage"),
    ("punch_state_display", "Etat du pointage"),
    ("verify_type_display", "Vérifier type"),
    ("work_code", "Code de travail"),
    ("terminal_sn", "N° série pointeuse"),
    ("terminal_alias", "Pointeuse"),
    ("area_alias", "Zone"),
    ("upload_time", "Heure de remontée"),
]
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
DAYS_FR = ["Lundi", "Mardi", "Mercredi", "Jeudi", "Vendredi", "Samedi", "Dimanche"]

WORK_DAYS = {int(d) for d in os.environ.get("WORK_DAYS", "0,1,2,3").split(",")}
DAILY_HOURS = float(os.environ.get("DAILY_HOURS", "9"))


def add_sheet(wb, title, headers, rows, warn_col=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
        if warn_col is not None and row[warn_col]:
            for cell in ws[ws.max_row]:
                cell.fill = WARN_FILL
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, header in enumerate(headers, 1):
        width = max([len(str(header))] + [len(str(r[i - 1])) for r in rows[:500]])
        ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 60)


def parse_date(value):
    return datetime.strptime(value, "%Y-%m-%d")


def hhmm(value):
    return f"{value:%H:%M}" if value else ""


def main():
    end = parse_date(sys.argv[2]) + timedelta(days=1, seconds=-1) if len(sys.argv) > 2 else zk.now_local()
    start = parse_date(sys.argv[1]) if len(sys.argv) > 1 else (end - timedelta(days=30)).replace(hour=0, minute=0, second=0)

    print(f"Lecture des pointages BioTime du {start} au {end}...")
    biotime = zk.BioTime()
    txns = sorted({t["id"]: t for t in biotime.transactions(start, end)}.values(),
                  key=lambda t: (str(t.get("emp_code")), t.get("punch_time") or ""))
    print(f"{len(txns)} pointages récupérés.")
    attendance, door, names = zk.group_punches(txns)

    try:
        employees = zk.Odoo().employee_map()
    except Exception as e:  # the preview works without Odoo
        print(f"Odoo indisponible, aperçu sans correspondance employés : {e}")
        employees = {}

    per_terminal = Counter(str(t.get("terminal_sn") or "") for t in txns)
    try:
        terminals = {str(d.get("sn")): d for d in biotime.terminals()}
    except Exception as e:
        print(f"Liste des pointeuses indisponible : {e}")
        terminals = {}
    usage = lambda sn: "Ignoré (ouverture porte)" if sn in zk.TERMINALS_IGNORE else "Pointage de présence"
    terminal_rows = [[sn, d.get("alias") or "", d.get("area_name") or "", d.get("ip_address") or "",
                      d.get("last_activity") or "", usage(sn), per_terminal.get(sn, 0)]
                     for sn, d in terminals.items()]
    terminal_rows += [[sn, "", "", "", "", usage(sn), n] for sn, n in per_terminal.items() if sn not in terminals]

    rows, anomalies, summary = [], [], []
    for code in sorted(set(attendance) | set(door), key=lambda c: (len(c), c)):
        emp = employees.get(code)
        odoo_name = emp["name"] if emp else ("INTROUVABLE" if employees else "")
        pro = bool(emp and emp["pro"])
        days = zk.build_days(attendance.get(code, {}), door.get(code, {}), pro)
        total, total_gap, with_anomaly = 0.0, 0.0, 0
        for day, d in days.items():
            expected = DAILY_HOURS if day.weekday() in WORK_DAYS else 0.0
            gap = round(d["hours"] - expected, 2) if d["check_out"] and not d["anomaly"] else ""
            total += d["hours"]
            total_gap += gap or 0.0
            with_anomaly += bool(d["anomaly"])
            rows.append([code, names[code], odoo_name, "oui" if pro else "", day.isoformat(), DAYS_FR[day.weekday()],
                         hhmm(d["check_in"]), hhmm(d["pause_start"]), hhmm(d["pause_end"]),
                         hhmm(d["check_out"]) or "en cours", d["outings_count"], d["outings_hours"],
                         d["hours"], expected, gap, d["detail"], d["anomaly"]])
            if d["anomaly"]:
                anomalies.append([code, names[code], day.isoformat(), DAYS_FR[day.weekday()], d["anomaly"], d["detail"]])
        summary.append([code, names[code], odoo_name, len(days), with_anomaly, round(total, 2), round(total_gap, 2)])

    wb = Workbook()
    wb.remove(wb.active)
    add_sheet(wb, "Présences Odoo (1 par jour)",
              ["Numéro d'employé", "Nom BioTime", "Employé Odoo", "Sorties pro", "Date", "Jour",
               "Arrivée (check_in)", "Début pause", "Fin pause", "Départ (check_out)",
               "Nb sorties", "Sorties déduites (h)", "Heures travaillées", "Heures prévues", "Écart",
               "Détail des pointages", "Anomalie"], rows, warn_col=16)
    add_sheet(wb, "Anomalies", ["Numéro d'employé", "Nom", "Date", "Jour", "Anomalie", "Détail des pointages"],
              anomalies)
    add_sheet(wb, "Résumé par employé",
              ["Numéro d'employé", "Nom BioTime", "Employé Odoo", "Jours", "Jours avec anomalie",
               "Heures travaillées", "Écart vs prévu (jours complets)"], summary, warn_col=4)
    add_sheet(wb, "Pointeuses", ["N° série", "Nom", "Zone", "Adresse IP", "Dernière activité",
                                 "Utilisation", "Nb pointages sur la période"], terminal_rows)
    add_sheet(wb, "Pointages bruts", [h for _, h in RAW_COLUMNS] + ["Utilisation"],
              [[t.get(k) if t.get(k) is not None else "" for k, _ in RAW_COLUMNS]
               + ["Ignoré (ouverture porte)" if zk.is_door_punch(t) else "Pointage de présence"] for t in txns])

    out = zk.BASE_DIR / f"apercu_biotime_{start:%Y%m%d}_{end:%Y%m%d}.xlsx"
    wb.save(out)
    print(f"Fichier généré : {out}")
    print(f"  {len(rows)} journée(s), {len(anomalies)} anomalie(s), {len(summary)} employé(s)")


if __name__ == "__main__":
    main()
